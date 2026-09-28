"""Text-layer PDF and DOCX extraction with explicit table-cell evidence."""

from collections import defaultdict
from io import BytesIO
from typing import NamedTuple

import pdfplumber
from docx import Document as WordDocument


class ExtractedChunk(NamedTuple):
    page: int
    text: str
    kind: str = "prose"
    table_id: str | None = None
    row_index: int | None = None
    column_index: int | None = None
    row_header: str = ""
    column_header: str = ""


def _prose_outside_tables(page, tables) -> str:
    boxes = [table.bbox for table in tables]
    lines: dict[int, list[tuple[float, str]]] = defaultdict(list)
    for char in page.chars:
        x = (char["x0"] + char["x1"]) / 2
        y = (char["top"] + char["bottom"]) / 2
        if any(x0 <= x <= x1 and top <= y <= bottom for x0, top, x1, bottom in boxes):
            continue
        lines[round(char["top"])].append((char["x0"], char["text"]))
    result: list[str] = []
    for line in sorted(lines):
        chars = sorted(lines[line])
        parts: list[str] = []
        prior_x = None
        for x, value in chars:
            if prior_x is not None and x - prior_x > 2.5 and parts and not parts[-1].endswith(" "):
                parts.append(" ")
            parts.append(value)
            prior_x = x
        text = "".join(parts).strip()
        if text:
            result.append(text)
    return "\n".join(result)


def extract_pdf(data: bytes) -> tuple[list[ExtractedChunk], int, list[int], list[int]]:
    chunks: list[ExtractedChunk] = []
    failed_pages: list[int] = []
    failed_tables: list[int] = []
    with pdfplumber.open(BytesIO(data)) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            try:
                tables = page.find_tables()
                for table_number, table in enumerate(tables, start=1):
                    grid = table.extract()
                    if not grid or len(grid) < 2 or not any(any(cell for cell in row) for row in grid):
                        failed_tables.append(page_number)
                        continue
                    headers = [str(value or "").strip() for value in grid[0]]
                    for row_index, row in enumerate(grid[1:], start=1):
                        row_values = [str(value or "").strip() for value in row]
                        row_header = row_values[0] if row_values else ""
                        for column_index, value in enumerate(row_values):
                            if not value:
                                continue
                            header = headers[column_index] if column_index < len(headers) else ""
                            evidence = f"Row: {row_header}; Column: {header}; Value: {value}"
                            chunks.append(
                                ExtractedChunk(
                                    page_number,
                                    evidence,
                                    "cell",
                                    f"p{page_number}-t{table_number}",
                                    row_index,
                                    column_index,
                                    row_header,
                                    header,
                                )
                            )
                prose = _prose_outside_tables(page, tables)
                if prose.strip():
                    chunks.extend(
                        ExtractedChunk(page_number, piece.strip())
                        for piece in prose.split("\n")
                        if piece.strip()
                    )
                elif not tables:
                    failed_pages.append(page_number)
            except Exception:
                failed_tables.append(page_number)
                prose = _prose_outside_tables(page, [])
                if prose.strip():
                    chunks.extend(ExtractedChunk(page_number, line) for line in prose.splitlines() if line.strip())
                else:
                    failed_pages.append(page_number)
    return chunks, len(pdf.pages), sorted(set(failed_pages)), sorted(set(failed_tables))


def extract_docx(data: bytes) -> tuple[list[ExtractedChunk], int, list[int], list[int]]:
    document = WordDocument(BytesIO(data))
    chunks: list[ExtractedChunk] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            chunks.append(ExtractedChunk(1, text))
    failed_tables: list[int] = []
    for table_number, table in enumerate(document.tables, start=1):
        grid = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if len(grid) < 2:
            failed_tables.append(1)
            continue
        headers = grid[0]
        for row_index, row in enumerate(grid[1:], start=1):
            row_header = row[0] if row else ""
            for column_index, value in enumerate(row):
                if value:
                    column_header = headers[column_index] if column_index < len(headers) else ""
                    chunks.append(
                        ExtractedChunk(
                            1,
                            f"Row: {row_header}; Column: {column_header}; Value: {value}",
                            "cell",
                            f"docx-t{table_number}",
                            row_index,
                            column_index,
                            row_header,
                            column_header,
                        )
                    )
    return chunks, 1, [], failed_tables


def extract_document(filename: str, data: bytes) -> tuple[list[ExtractedChunk], int, list[int], list[int]]:
    if filename.lower().endswith(".pdf"):
        return extract_pdf(data)
    if filename.lower().endswith(".docx"):
        return extract_docx(data)
    raise ValueError("Unsupported document format")
